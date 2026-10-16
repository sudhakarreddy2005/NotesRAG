import os
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, List
from langchain_core.documents import Document
from langchain_community.document_loaders import (
    CSVLoader,
    JSONLoader,
    TextLoader,
)

BINARY_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt",
    ".zip", ".tar", ".gz", ".7z", ".rar", ".bz2",
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico", ".tiff",
    ".mp3", ".wav", ".ogg", ".mp4", ".avi", ".mkv", ".mov",
    ".bin", ".exe", ".dll", ".so", ".dylib", ".iso", ".pyc"
}


def is_binary_file(path: Path) -> bool:
    """Detect if a file is binary by extension or null-byte inspection."""
    if path.suffix.lower() in BINARY_EXTENSIONS:
        return True
    try:
        with open(path, "rb") as f:
            chunk = f.read(8192)
            if b"\x00" in chunk:
                return True
    except Exception:
        pass
    return False


def load_docx_document(path: Path) -> List[Document]:
    """Robust multi-tier DOCX loader:
    1. Try Docx2txtLoader (LangChain)
    2. Try python-docx
    3. Built-in zero-dependency ZipFile + XML parser
    """
    # 1. Docx2txtLoader
    try:
        from langchain_community.document_loaders import Docx2txtLoader
        loader = Docx2txtLoader(str(path))
        docs = loader.load()
        if docs and any(d.page_content.strip() for d in docs):
            return docs
    except Exception as e:
        print(f"[DEBUG] Docx2txtLoader failed for {path.name}: {e}")

    # 2. python-docx
    try:
        import docx
        doc = docx.Document(str(path))
        lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        full_text = "\n\n".join(lines)
        if full_text.strip():
            return [Document(page_content=full_text, metadata={"source": str(path)})]
    except Exception as e:
        print(f"[DEBUG] python-docx failed for {path.name}: {e}")

    # 3. Built-in standard library fallback
    try:
        with zipfile.ZipFile(str(path)) as z:
            if "word/document.xml" in z.namelist():
                xml_data = z.read("word/document.xml")
                root = ET.fromstring(xml_data)
                paras = []
                for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                    texts = [n.text for n in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if n.text]
                    if texts:
                        paras.append("".join(texts).strip())
                full_text = "\n\n".join(p for p in paras if p)
                if full_text.strip():
                    return [Document(page_content=full_text, metadata={"source": str(path)})]
    except Exception as e:
        print(f"[ERROR] Built-in zipfile DOCX parser failed for {path.name}: {e}")

    raise ValueError(f"Unable to extract readable text from Word document: {path.name}")


def load_pdf_document(path: Path) -> List[Document]:
    """Robust PDF loader with PyMuPDF and PyPDF fallback."""
    try:
        from langchain_community.document_loaders import PyMuPDFLoader
        loader = PyMuPDFLoader(str(path))
        docs = loader.load()
        if docs:
            return docs
    except Exception:
        pass

    try:
        from langchain_community.document_loaders import PyPDFLoader
        loader = PyPDFLoader(str(path))
        docs = loader.load()
        if docs:
            return docs
    except Exception as e:
        print(f"[ERROR] Failed to load PDF {path.name}: {e}")

    return []


def load_xlsx_document(path: Path) -> List[Document]:
    """Extract text from Excel spreadsheets (.xlsx) using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        lines = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines.append(f"=== Sheet: {sheet} ===")
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    row_str = " | ".join(str(cell).strip() if cell is not None else "" for cell in row)
                    lines.append(row_str)
        text = "\n".join(lines)
        if text.strip():
            return [Document(page_content=text, metadata={"source": str(path)})]
    except Exception as e:
        print(f"[ERROR] Failed to load Excel file {path.name}: {e}")
    return []


def load_all_documents(data_dir: str = "data") -> List[Any]:
    """Load all supported documents from the data directory into LangChain Document format.

    Supported formats: .pdf, .txt, .docx, .json, .csv, .xlsx
    """
    data_path = Path(data_dir).resolve()
    print(f"[INFO] Scanning data path: {data_path}")

    if not data_path.exists():
        print(f"[WARNING] Data directory not found: {data_path}")
        return []

    documents: List[Any] = []

    # PDF documents
    pdf_files = list(data_path.glob("**/*.pdf"))
    if pdf_files:
        print(f"[INFO] Found {len(pdf_files)} PDF file(s).")
    for pdf_file in pdf_files:
        try:
            loaded = load_pdf_document(pdf_file)
            for doc in loaded:
                doc.metadata["source_name"] = pdf_file.name
                doc.metadata["file_type"] = "pdf"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded {len(loaded)} page(s) from {pdf_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load PDF file {pdf_file.name}: {e}")

    # Plain text documents
    text_files = list(data_path.glob("**/*.txt"))
    if text_files:
        print(f"[INFO] Found {len(text_files)} TXT file(s).")
    for text_file in text_files:
        try:
            loader = TextLoader(str(text_file), encoding="utf-8")
            loaded = loader.load()
            for doc in loaded:
                doc.metadata["source_name"] = text_file.name
                doc.metadata["file_type"] = "txt"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded document from {text_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load TXT file {text_file.name}: {e}")

    # Microsoft Word documents (.docx)
    docx_files = list(data_path.glob("**/*.docx"))
    if docx_files:
        print(f"[INFO] Found {len(docx_files)} DOCX file(s).")
    for docx_file in docx_files:
        try:
            loaded = load_docx_document(docx_file)
            for doc in loaded:
                doc.metadata["source_name"] = docx_file.name
                doc.metadata["file_type"] = "docx"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded document from {docx_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load DOCX file {docx_file.name}: {e}")

    # JSON documents
    json_files = list(data_path.glob("**/*.json"))
    if json_files:
        print(f"[INFO] Found {len(json_files)} JSON file(s).")
    for json_file in json_files:
        try:
            loader = JSONLoader(str(json_file), jq_schema=".", text_content=False)
            loaded = loader.load()
            for doc in loaded:
                doc.metadata["source_name"] = json_file.name
                doc.metadata["file_type"] = "json"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded JSON document from {json_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load JSON file {json_file.name}: {e}")

    # CSV documents
    csv_files = list(data_path.glob("**/*.csv"))
    if csv_files:
        print(f"[INFO] Found {len(csv_files)} CSV file(s).")
    for csv_file in csv_files:
        try:
            loader = CSVLoader(str(csv_file))
            loaded = loader.load()
            for doc in loaded:
                doc.metadata["source_name"] = csv_file.name
                doc.metadata["file_type"] = "csv"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded CSV document from {csv_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load CSV file {csv_file.name}: {e}")

    # Excel documents (.xlsx)
    xlsx_files = list(data_path.glob("**/*.xlsx"))
    if xlsx_files:
        print(f"[INFO] Found {len(xlsx_files)} XLSX file(s).")
    for xlsx_file in xlsx_files:
        try:
            loaded = load_xlsx_document(xlsx_file)
            for doc in loaded:
                doc.metadata["source_name"] = xlsx_file.name
                doc.metadata["file_type"] = "xlsx"
            documents.extend(loaded)
            print(f"[DEBUG] Loaded Excel document from {xlsx_file.name}")
        except Exception as e:
            print(f"[ERROR] Failed to load XLSX file {xlsx_file.name}: {e}")

    print(f"[INFO] Ingestion complete. Total document chunks loaded: {len(documents)}")
    return documents


def load_single_document(file_path: str) -> List[Any]:
    """Load any supported document kind from file_path into LangChain Document format.
    Guarantees that binary archives and files are never loaded as raw text bytes.
    """
    path = Path(file_path).resolve()
    ext = path.suffix.lower()
    documents: List[Any] = []

    try:
        if ext == ".pdf":
            documents = load_pdf_document(path)
        elif ext in [".txt", ".md", ".text", ".py", ".js", ".ts", ".html", ".css", ".yaml", ".yml", ".toml", ".sh"]:
            loader = TextLoader(str(path), encoding="utf-8")
            documents = loader.load()
        elif ext == ".docx":
            documents = load_docx_document(path)
        elif ext in [".xlsx", ".xls"]:
            documents = load_xlsx_document(path)
        elif ext == ".json":
            loader = JSONLoader(str(path), jq_schema=".", text_content=False)
            documents = loader.load()
        elif ext == ".csv":
            loader = CSVLoader(str(path))
            documents = loader.load()
        else:
            # Strict safety guard: never open binary files as text
            if is_binary_file(path):
                print(f"[ERROR] Refusing to treat binary file as plain text: {path.name}")
                return []

            # Universal text fallback for other plain-text files
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            if content.strip():
                documents = [Document(page_content=content, metadata={"source": str(path)})]
    except Exception as e:
        print(f"[ERROR] Failed to load {path.name}: {e}")
        return []

    for doc in documents:
        doc.metadata["source_name"] = path.name
        doc.metadata["file_type"] = ext.lstrip(".") or "unknown"

    print(f"[INFO] Loaded {len(documents)} document chunk(s) from '{path.name}'.")
    return documents